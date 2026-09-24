"""Funcional Lumos: servidor web, rotas da API e sincronização em segundo plano."""
import base64
import io
import json
import os
import re
import secrets
import threading
import time
import traceback
from datetime import datetime, timedelta
from pathlib import Path

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from sqlalchemy import func

from . import ai, mail
from .db import (Analysis, Decision, Document, Email, SessionLocal, get_setting,
                 init_db)
from .mail import index_email, normalize_subject, strip_quoted
from .retrieval import get_index, mark_dirty, rank_texts

APP_USER = os.getenv("APP_USER", "startgi")
APP_PASSWORD = os.getenv("APP_PASSWORD", "")
POLL_MINUTES = int(os.getenv("POLL_MINUTES", "5"))
AUTO_ANALYZE = os.getenv("AUTO_ANALYZE", "true").lower() == "true"
STATIC = Path(__file__).parent / "static"

app = FastAPI(title="Funcional Lumos")
app.mount("/static", StaticFiles(directory=STATIC), name="static")
state = {"last_sync": None, "last_error": None, "running": False, "log": []}


def log(msg):
    line = f"{datetime.now():%d/%m %H:%M:%S} {msg}"
    print(line, flush=True)
    state["log"] = (state["log"] + [line])[-60:]


@app.middleware("http")
async def basic_auth(request: Request, call_next):
    if request.url.path == "/healthz" or not APP_PASSWORD:
        return await call_next(request)
    header = request.headers.get("authorization", "")
    ok = False
    if header.lower().startswith("basic "):
        try:
            user, _, pwd = base64.b64decode(header[6:]).decode().partition(":")
            ok = secrets.compare_digest(user, APP_USER) and secrets.compare_digest(pwd, APP_PASSWORD)
        except Exception:
            ok = False
    if not ok:
        return Response(status_code=401, headers={"WWW-Authenticate": 'Basic realm="Funcional Lumos"'})
    return await call_next(request)


def humanize_error(ex):
    if isinstance(ex, OSError) and not isinstance(ex, __import__("imaplib").IMAP4.error):
        return (f"Não foi possível alcançar o servidor {mail.IMAP_HOST}. Confira o ZOHO_IMAP_HOST "
                f"({ex.__class__.__name__}).")
    msg = ex.args[0] if getattr(ex, "args", None) else ex
    if isinstance(msg, bytes):
        msg = msg.decode(errors="replace")
    msg = str(msg)
    low = msg.lower()
    if "enable imap" in low:
        return ("O IMAP não está habilitado nesta conta do Zoho. Habilite em Zoho Mail → Configurações → "
                "Contas de e-mail → Acesso IMAP. Se for conta de organização, o administrador também precisa "
                "liberar o IMAP no Admin Console.")
    if "authentication" in low or "invalid credentials" in low or "login" in low and "fail" in low:
        return ("O Zoho recusou o login. Confira o ZOHO_EMAIL e use uma senha específica de aplicativo "
                "no ZOHO_APP_PASSWORD, não a senha normal.")
    if "getaddrinfo" in low or "timed out" in low or "connection refused" in low:
        return f"Não foi possível alcançar o servidor {mail.IMAP_HOST}. Confira o ZOHO_IMAP_HOST."
    return msg


def run_sync():
    if state["running"]:
        return
    state["running"] = True
    try:
        new_ids = mail.sync(log=log)
        state["last_sync"] = datetime.now().isoformat(timespec="seconds")
        state["last_error"] = None
        log(f"Sincronização concluída. Novos e-mails da Motiva: {len(new_ids)}")
        if AUTO_ANALYZE and ai.API_KEY:
            for eid in new_ids:
                try:
                    ai.analyze(eid)
                    log(f"Briefing gerado para o e-mail {eid}")
                except Exception as ex:
                    log(f"Falha ao analisar o e-mail {eid}: {ex}")
    except Exception as ex:
        state["last_error"] = humanize_error(ex)
        log(f"Erro na sincronização: {state['last_error']}")
    finally:
        state["running"] = False


def scheduler():
    while True:
        if mail.configured():
            run_sync()
        time.sleep(POLL_MINUTES * 60)


@app.on_event("startup")
def startup():
    init_db()
    threading.Thread(target=scheduler, daemon=True).start()


@app.get("/healthz")
def healthz():
    return {"ok": True}


@app.get("/")
def index():
    return FileResponse(STATIC / "index.html")


# ---------- dashboard ----------
@app.get("/api/dashboard")
def dashboard():
    with SessionLocal() as s:
        emails_total = s.query(Email).count()
        emails_in = s.query(Email).filter(Email.direction == "in").count()
        emails_out = s.query(Email).filter(Email.direction == "out").count()
        emails_novos = s.query(Email).filter(Email.direction == "in", Email.status == "novo").count()
        emails_analisados = s.query(Email).filter(Email.direction == "in", Email.status == "analisado").count()
        emails_respondidos = s.query(Email).filter(Email.direction == "in", Email.status == "respondido").count()
        emails_ignorados = s.query(Email).filter(Email.direction == "in", Email.status == "ignorado").count()
        emails_arquivados = s.query(Email).filter(Email.status == "arquivado").count()
        decisoes_total = s.query(Decision).count()
        decisoes_vigentes = s.query(Decision).filter(Decision.status == "vigente").count()
        decisoes_proposta = s.query(Decision).filter(Decision.status == "proposta").count()
        decisoes_discussao = s.query(Decision).filter(Decision.status == "em_discussao").count()
        decisoes_substituida = s.query(Decision).filter(Decision.status == "substituida").count()
        decisoes_descartada = s.query(Decision).filter(Decision.status == "descartada").count()
        documentos = s.query(Document).count()
        analyses_total = s.query(Analysis).count()
        conversas = s.query(Email.thread_key).distinct().count()
        last_sync = get_setting(s, "last_sync")
        first_email_date = s.query(Email).order_by(Email.date.asc()).first()
        last_email_date = s.query(Email).order_by(Email.date.desc()).first()
        ed = first_email_date.date.isoformat() if first_email_date and first_email_date.date else None
        ld = last_email_date.date.isoformat() if last_email_date and last_email_date.date else None
        days = 21
        since = (datetime.utcnow() - timedelta(days=days - 1)).date()
        dia_total = {}
        for dia, total in s.query(func.date(Email.date).label("dia"), func.count(Email.id)) \
                .filter(Email.date >= since, Email.date.isnot(None)).group_by("dia").all():
            dia_total[str(dia)] = total
        por_dia = [{"dia": d.isoformat(), "total": dia_total.get(d.isoformat(), 0)}
                   for d in (since + timedelta(days=i) for i in range(days))]
        pend = s.query(Email).filter(Email.direction == "in",
                                     Email.status.in_(["novo", "analisado"])).order_by(Email.date.desc()).limit(400).all()
        resp = {}
        for e in pend:
            t = resp.setdefault(e.thread_key, {"thread_key": e.thread_key, "subject": e.subject,
                                               "count": 0, "analyzed": 0, "last_date": e.date,
                                               "last_from": e.from_addr})
            t["count"] += 1
            if e.status == "analisado":
                t["analyzed"] += 1
        para_responder = sorted(resp.values(), key=lambda t: t["last_date"] or datetime.min, reverse=True)[:8]
        for t in para_responder:
            t["last_date"] = t["last_date"].isoformat() if t["last_date"] else None
    zoho_state = "nao_configurado" if not mail.configured() else ("erro" if state["last_error"] else ("ok" if last_sync else "aguardando"))
    log_count = len(state.get("log", []))
    return {
        "emails": {
            "total": emails_total, "in": emails_in, "out": emails_out,
            "novos": emails_novos, "analisados": emails_analisados,
            "respondidos": emails_respondidos, "ignorados": emails_ignorados,
            "arquivados": emails_arquivados,
        },
        "decisoes": {
            "total": decisoes_total, "vigentes": decisoes_vigentes,
            "proposta": decisoes_proposta, "em_discussao": decisoes_discussao,
            "substituida": decisoes_substituida, "descartada": decisoes_descartada,
        },
        "documentos": documentos,
        "analyses": analyses_total,
        "conversas": conversas,
        "por_dia": por_dia,
        "para_responder": para_responder,
        "ultima_sync": last_sync,
        "zoho_state": zoho_state,
        "ia": bool(ai.API_KEY),
        "modelo": f"{ai.PROVIDER}:{ai.MODEL}",
        "log_count": log_count,
        "log": state.get("log", [])[-20:],
        "first_email_date": ed,
        "last_email_date": ld,
    }


# ---------- status ----------
@app.get("/api/status")
def status():
    with SessionLocal() as s:
        counts = {
            "novos": s.query(Email).filter(Email.direction == "in", Email.status == "novo").count(),
            "analisados": s.query(Email).filter(Email.direction == "in", Email.status == "analisado").count(),
            "emails": s.query(Email).count(),
            "decisoes_vigentes": s.query(Decision).filter(Decision.status == "vigente").count(),
            "decisoes_propostas": s.query(Decision).filter(Decision.status == "proposta").count(),
            "documentos": s.query(Document).count(),
        }
        last_db = get_setting(s, "last_sync")
    zoho_state = "nao_configurado" if not mail.configured() else ("erro" if state["last_error"] else ("ok" if last_db else "aguardando"))
    return {"zoho": mail.configured(), "zoho_state": zoho_state, "zoho_user": mail.IMAP_USER, "ia": bool(ai.API_KEY),
            "modelo": f"{ai.PROVIDER}:{ai.MODEL}", "dominios_motiva": mail.MOTIVA_DOMAINS,
            "poll_minutos": POLL_MINUTES, "auto_analise": AUTO_ANALYZE, "ultima_sync": last_db,
            "sincronizando": state["running"], "erro": state["last_error"], "log": state["log"][-25:], **counts}


@app.post("/api/sync")
def sync_now():
    if not mail.configured():
        raise HTTPException(400, "Configure ZOHO_EMAIL e ZOHO_APP_PASSWORD para sincronizar.")
    threading.Thread(target=run_sync, daemon=True).start()
    return {"ok": True}


# ---------- caixa ----------
def make_snippet(text, query, width=170):
    """Trecho do texto ao redor da primeira palavra da busca, para confirmar o hit sem abrir o e-mail."""
    t = re.sub(r"\s+", " ", text or "").strip()
    if not t:
        return ""
    words = [w for w in query.split() if len(w) > 2]
    pos = min((t.lower().find(w.lower()) for w in words if t.lower().find(w.lower()) >= 0), default=-1)
    if pos < 0:
        return t[:width] + ("…" if len(t) > width else "")
    start = max(0, pos - 50)
    end = min(len(t), start + width)
    return ("…" if start else "") + t[start:end] + ("…" if end < len(t) else "")


@app.get("/api/threads")
def threads(filtro: str = "todas", q: str = "", data_inicio: str = "", data_fim: str = "", remetente: str = ""):
    with SessionLocal() as s:
        query = s.query(Email).order_by(Email.date.desc())
        if data_inicio:
            try: query = query.filter(Email.date >= datetime.fromisoformat(data_inicio))
            except Exception: pass
        if data_fim:
            try: query = query.filter(Email.date <= datetime.fromisoformat(data_fim + "T23:59:59"))
            except Exception: pass
        if remetente:
            query = query.filter(Email.from_addr.ilike(f"%{remetente}%"))
        rows = query.limit(3000).all()
    by = {}
    for e in rows:
        t = by.setdefault(e.thread_key, {"thread_key": e.thread_key, "subject": e.subject, "last_date": e.date,
                                          "count": 0, "pending": 0, "analyzed": 0, "last_from": None,
                                          "last_in_id": None, "senders": set(), "subjects": set(),
                                          "last_dir": e.direction,
                                          "st": {"novo": 0, "analisado": 0, "respondido": 0,
                                                 "arquivado": 0, "ignorado": 0},
                                          "preview": ""})
        t["count"] += 1
        if e.from_addr:
            t["senders"].add(e.from_addr)
        if e.subject:
            t["subjects"].add(e.subject)
        if e.direction == "in":
            t["st"][e.status] = t["st"].get(e.status, 0) + 1
            if e.status in ("novo", "analisado"):
                t["pending"] += 1
                if e.status == "analisado":
                    t["analyzed"] += 1
            if t["last_in_id"] is None:
                t["last_in_id"] = e.id
        if t["last_from"] is None:
            t["last_from"] = "StartGi" if e.direction == "out" else e.from_addr
        if not t["preview"]:
            t["preview"] = re.sub(r"\s+", " ", (e.body_clean or e.body or "")).strip()[:170]
    with SessionLocal() as s2:
        in_ids = [t["last_in_id"] for t in by.values() if t["last_in_id"]]
        cls = {}
        if in_ids:
            for a in s2.query(Analysis).filter(Analysis.email_id.in_(in_ids)).order_by(Analysis.id).all():
                try:
                    cls[a.email_id] = (json.loads(a.briefing_json or "{}").get("classificacao") or {}).get("tipo")
                except Exception:
                    pass
    for t in by.values():
        t["classificacao"] = cls.get(t["last_in_id"])
        st = t["st"]
        if st["novo"]:
            t["status"] = "novo"
        elif st["analisado"]:
            t["status"] = "pronto"
        elif st["respondido"] or t["last_dir"] == "out":
            t["status"] = "aguardando"  # StartGi já respondeu, aguardando o retorno do cliente
        else:
            t["status"] = "arquivado"
    all_threads = list(by.values())
    counts = {
        "pendentes": sum(1 for t in all_threads if t["status"] in ("novo", "pronto")),
        "aguardando": sum(1 for t in all_threads if t["status"] == "aguardando"),
    }
    out = all_threads
    ql = (q or "").strip()
    if not ql:  # a busca é global: com termo, varre todas as conversas
        if filtro == "pendentes":
            out = [t for t in out if t["status"] in ("novo", "pronto")]
        elif filtro == "aguardando":
            out = [t for t in out if t["status"] == "aguardando"]
    snippet = {}
    if ql:
        ql_low = ql.lower()
        matched = {t["thread_key"] for t in out
                   if any(ql_low in (sub or "").lower() for sub in t["subjects"])
                   or any(ql_low in snd.lower() for snd in t["senders"])}
        try:
            idx = get_index()
            hits = idx.search(ql, k=40, kinds={"email"}) if idx else []
        except Exception:
            hits = []
        email_ids = list({sid for _, (_, _, sid, _, _) in hits})
        id2key = {}
        if email_ids:
            with SessionLocal() as s3:
                id2key = dict(s3.query(Email.id, Email.thread_key).filter(Email.id.in_(email_ids)).all())
        for sc, (cid, kind, sid, label, text) in hits:
            key = id2key.get(sid)
            if key and key in by:
                matched.add(key)
                if key not in snippet:
                    snippet[key] = make_snippet(text, ql)
        out = [t for t in out if t["thread_key"] in matched]
    prio = {"novo": 0, "pronto": 1, "aguardando": 2, "arquivado": 3}
    out.sort(key=lambda t: t["last_date"] or datetime.min, reverse=True)
    out.sort(key=lambda t: prio.get(t["status"], 9))
    return {"counts": counts, "items": [{
        "thread_key": t["thread_key"], "subject": t["subject"], "count": t["count"],
        "pending": t["pending"], "analyzed": t["analyzed"], "last_from": t["last_from"],
        "last_date": t["last_date"].isoformat() if t["last_date"] else None,
        "classificacao": t["classificacao"], "status": t["status"],
        "preview": t["preview"], "snippet": snippet.get(t["thread_key"]),
    } for t in out[:400]]}


def email_dict(e, full=True):
    d = {"id": e.id, "direction": e.direction, "from": e.from_addr, "to": e.to_addr, "cc": e.cc_addr,
         "subject": e.subject, "date": e.date.isoformat() if e.date else None, "status": e.status}
    if full:
        d["body"] = e.body_clean or e.body
        d["body_full"] = e.body
    return d


def analysis_dict(a):
    return {"id": a.id, "email_id": a.email_id, "briefing": json.loads(a.briefing_json or "{}"),
            "draft_subject": a.draft_subject, "draft_body": a.draft_body, "final_body": a.final_body,
            "model": a.model, "sources": json.loads(a.sources_json or "[]"),
            "created_at": a.created_at.isoformat() if a.created_at else None}


@app.get("/api/threads/{thread_key:path}")
def thread_detail(thread_key: str):
    with SessionLocal() as s:
        emails = s.query(Email).filter(Email.thread_key == thread_key).order_by(Email.date).all()
        if not emails:
            raise HTTPException(404, "Conversa não encontrada.")
        ids = [e.id for e in emails]
        analyses = s.query(Analysis).filter(Analysis.email_id.in_(ids)).order_by(Analysis.id).all()
        return {"emails": [email_dict(e) for e in emails], "analyses": [analysis_dict(a) for a in analyses]}


@app.post("/api/emails/{email_id}/analyze")
def analyze_email(email_id: int):
    try:
        aid = ai.analyze(email_id)
    except Exception as ex:
        traceback.print_exc()
        raise HTTPException(500, f"Não foi possível gerar o briefing: {ex}")
    with SessionLocal() as s:
        return analysis_dict(s.get(Analysis, aid))


class RewriteIn(BaseModel):
    instrucao: str
    corpo_atual: str | None = None


@app.post("/api/analyses/{analysis_id}/rewrite")
def rewrite(analysis_id: int, body: RewriteIn):
    try:
        return {"draft_body": ai.rewrite(analysis_id, body.instrucao, body.corpo_atual)}
    except Exception as ex:
        raise HTTPException(500, f"Não foi possível reescrever: {ex}")


class RespondIn(BaseModel):
    corpo_final: str


@app.post("/api/analyses/{analysis_id}/sent")
def mark_sent(analysis_id: int, body: RespondIn):
    with SessionLocal() as s:
        a = s.get(Analysis, analysis_id)
        if not a:
            raise HTTPException(404, "Análise não encontrada.")
        a.final_body = body.corpo_final
        e = a.email
        s.query(Email).filter(Email.thread_key == e.thread_key, Email.direction == "in",
                              Email.status.in_(["novo", "analisado"])).update({"status": "respondido"},
                                                                              synchronize_session=False)
        s.commit()
    return {"ok": True}


@app.post("/api/emails/{email_id}/status")
def set_email_status(email_id: int, status: str = Form(...)):
    if status not in ("novo", "analisado", "respondido", "ignorado", "arquivado"):
        raise HTTPException(400, "Status inválido.")
    with SessionLocal() as s:
        e = s.get(Email, email_id)
        e.status = status
        s.commit()
    return {"ok": True}


class ManualEmailIn(BaseModel):
    remetente: str
    assunto: str
    corpo: str
    direcao: str = "in"


@app.post("/api/emails/manual")
def manual_email(body: ManualEmailIn):
    """Cola um e-mail manualmente (útil antes de conectar o Zoho, ou para mensagens de WhatsApp)."""
    with SessionLocal() as s:
        now = datetime.utcnow()
        key = normalize_subject(body.assunto)
        e = Email(message_id=f"manual-{now.timestamp()}", thread_key=key, folder="manual",
                  direction="out" if body.direcao == "out" else "in", from_addr=body.remetente,
                  to_addr="", cc_addr="", subject=body.assunto, date=now, body=body.corpo,
                  body_clean=strip_quoted(body.corpo) or body.corpo, is_motiva=True,
                  status="novo" if body.direcao != "out" else "respondido")
        s.add(e)
        s.flush()
        index_email(s, e)
        s.commit()
        mark_dirty()
        return {"id": e.id, "thread_key": key}


# ---------- decisões ----------
class DecisionIn(BaseModel):
    titulo: str
    modulo: str = ""
    regra: str
    fonte: str = ""
    data_decisao: str = ""
    aprovado_por: str = ""
    status: str = "vigente"
    substituida_por: int | None = None
    notas: str = ""


def decision_dict(d):
    return {c: getattr(d, c) for c in ("id", "titulo", "modulo", "regra", "fonte", "data_decisao", "aprovado_por",
                                       "status", "substituida_por", "notas", "origem_email_id")} | {
        "updated_at": d.updated_at.isoformat() if d.updated_at else None}


@app.get("/api/decisions")
def list_decisions(status: str = ""):
    with SessionLocal() as s:
        q = s.query(Decision)
        if status:
            q = q.filter(Decision.status == status)
        return [decision_dict(d) for d in q.order_by(Decision.modulo, Decision.id.desc()).all()]


@app.post("/api/decisions")
def create_decision(body: DecisionIn):
    with SessionLocal() as s:
        d = Decision(**body.model_dump())
        s.add(d)
        s.commit()
        return decision_dict(d)


@app.put("/api/decisions/{decision_id}")
def update_decision(decision_id: int, body: DecisionIn):
    with SessionLocal() as s:
        d = s.get(Decision, decision_id)
        if not d:
            raise HTTPException(404, "Decisão não encontrada.")
        for k, v in body.model_dump().items():
            setattr(d, k, v)
        if body.substituida_por:
            d.status = "substituida"
        s.commit()
        return decision_dict(d)


# ---------- base de conhecimento ----------
def read_upload(name, raw):
    low = name.lower()
    if low.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(raw)).pages)
    if low.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text for c in r.cells))
        return "\n".join(parts)
    return raw.decode("utf-8", errors="replace")


@app.get("/api/documents")
def list_documents():
    with SessionLocal() as s:
        return [{"id": d.id, "nome": d.nome, "tipo": d.tipo, "tamanho": len(d.content or ""),
                 "created_at": d.created_at.isoformat() if d.created_at else None}
                for d in s.query(Document).order_by(Document.id.desc()).all()]


@app.post("/api/documents")
async def upload_document(tipo: str = Form("documentacao"), nome: str = Form(""), texto: str = Form(""),
                          arquivo: UploadFile | None = File(None)):
    content, filename = texto, nome
    if arquivo is not None and arquivo.filename:
        raw = await arquivo.read()
        try:
            content = read_upload(arquivo.filename, raw)
        except Exception as ex:
            raise HTTPException(400, f"Não consegui ler o arquivo: {ex}")
        filename = nome or arquivo.filename
    if not (content or "").strip():
        raise HTTPException(400, "Envie um arquivo ou cole o texto.")
    with SessionLocal() as s:
        d = Document(nome=filename or "Sem nome", tipo=tipo, content=content)
        s.add(d)
        s.flush()
        ai.index_document(s, d)
        s.commit()
        return {"id": d.id, "nome": d.nome, "tamanho": len(content)}


@app.delete("/api/documents/{doc_id}")
def delete_document(doc_id: int):
    from .db import Chunk
    with SessionLocal() as s:
        s.query(Chunk).filter(Chunk.source_kind == "document", Chunk.source_id == doc_id).delete()
        d = s.get(Document, doc_id)
        if d:
            s.delete(d)
        s.commit()
    mark_dirty()
    return {"ok": True}


# ---------- assistente (pergunta sobre todo o contexto) ----------
CHAT_SYSTEM = """Você é o assistente do Funcional Lumos da StartGi (projeto Lumos, cliente Motiva).
Você responde perguntas usando TODAS as fontes abaixo — e-mails históricos, decisões registradas e documentos
da base de conhecimento. Regras:

1. Responda em português, claro e direto, focado no que foi perguntado.
2. Cite as fontes entre colchetes sempre que usar uma delas: [E#] e-mail, [D#] decisão, [DOC#] documento.
   Se a resposta se basear em mais de uma fonte, cite todas.
3. Nunca invente. Se não houver fonte sobre o assunto, diga claramente que não há informação registrada.
4. Separe o que é fato documentado do que é inferência sua.
5. Se a pergunta não tiver relação com o projeto (Lumos, Motiva, StartGi, supply chain), responda que só
   trata de assuntos do projeto."""


class ChatIn(BaseModel):
    pergunta: str


@app.post("/api/chat")
def chat(body: ChatIn):
    if not ai.API_KEY:
        raise HTTPException(400, "IA não configurada: defina AI_API_KEY nas variáveis do Render.")
    pergunta = body.pergunta.strip()
    if not pergunta:
        raise HTTPException(400, "Escreva uma pergunta.")
    with SessionLocal() as s:
        idx = get_index()
        docs = idx.search(pergunta, k=8, kinds={"document"}) if idx else []
        mails = idx.search(pergunta, k=12, kinds={"email"}) if idx else []
        decisions = s.query(Decision).filter(Decision.status.in_(["vigente", "em_discussao", "proposta"])).all()
        ranked = rank_texts(pergunta, decisions, key=lambda d: f"{d.titulo} {d.modulo} {d.regra}")
        chosen = [d for sc, d in ranked if sc > 0][:25] or [d for d in decisions if d.status == "vigente"][:10]

    parts, sources = [], []
    parts.append("## Documentos e especificações da base")
    if not docs:
        parts.append("(nenhum documento relacionado)")
    for sc, (cid, kind, sid, label, text) in docs:
        parts.append(f"[DOC{sid}] {label}\n{text}\n")
        sources.append({"code": f"DOC{sid}", "label": label})
    parts.append("## E-mails do histórico (ordem de relevância)")
    if not mails:
        parts.append("(nenhum e-mail relacionado)")
    for sc, (cid, kind, sid, label, text) in mails:
        parts.append(f"[E{sid}] {label}\n{text}\n")
        sources.append({"code": f"E{sid}", "label": label})
    parts.append("## Registro de decisões")
    if not chosen:
        parts.append("(nenhuma decisão registrada relacionada)")
    for d in chosen:
        parts.append(f"[D{d.id}] {d.status.upper()} | {d.modulo} | {d.titulo}\nRegra: {d.regra}\nFonte: {d.fonte} | {d.data_decisao}")
        sources.append({"code": f"D{d.id}", "label": d.titulo})

    user = f"{chr(10).join(parts)}\n\n## Pergunta do usuário\n{pergunta}"
    raw = ai.call_model(CHAT_SYSTEM, user)

    seen, uniq = set(), []
    for src in sources:
        if src["code"] not in seen:
            seen.add(src["code"])
            uniq.append(src)
    return {"resposta": raw.strip(), "fontes": uniq}
