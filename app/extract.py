"""Extrai texto de arquivos binários e de texto, usado por uploads e anexos de e-mail."""
import io


def extract_text_from_bytes(name, raw):
    low = (name or "").lower()
    if low.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(raw)).pages)
    if low.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text for c in r.cells))
        return "\n".join(parts)
    if low.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        parts = []
        try:
            for ws in wb.worksheets:
                parts.append(f"Planilha: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        finally:
            wb.close()
        return "\n".join(parts)
    if low.endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(file_contents=raw)
        parts = []
        for ws in wb.sheets():
            parts.append(f"Planilha: {ws.name}")
            for i in range(ws.nrows):
                cells = [str(c).strip() for c in ws.row_values(i) if str(c).strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if low.endswith((".txt", ".md", ".csv", ".json", ".jsonl")):
        return raw.decode("utf-8", errors="replace")
    return ""